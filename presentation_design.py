"""Deterministic slide scenes shared by HTML previews and editable PPTX output."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def _units(text):
    return sum(1 if ord(c) > 255 else .54 for c in str(text))


def _wrap(text, capacity):
    lines, line, used = [], '', 0
    for char in str(text):
        width = 1 if ord(char) > 255 else .54
        if char == '\n' or used + width > capacity:
            lines.append(line)
            line, used = '', 0
            if char == '\n':
                continue
        line += char
        used += width
    lines.append(line)
    return '\n'.join(lines)


def _text(scene, text, x, y, w, h, size=22, color='#232824', bold=False, font='Microsoft YaHei'):
    # Explicit line breaks give the browser and PowerPoint the same line lengths.
    for fitted in range(size, 7, -1):
        wrapped = _wrap(text, max(1, w * 72 / fitted))
        if len(wrapped.split('\n')) * fitted * 1.22 <= h * 72:
            break
    scene['elements'].append(dict(kind='text', text=wrapped, x=x, y=y, w=w, h=h,
                                  size=fitted, color=color, bold=bold, font=font))


def _rect(scene, x, y, w, h, color, kind='rect'):
    scene['elements'].append(dict(kind=kind, x=x, y=y, w=w, h=h, color=color))


def _assign_layouts(expanded, natural_layout, threshold=0.70):
    """为每个 scene 分配最终布局；内容页做「版面重复率」约束（单一布局占比 ≤ threshold）。"""
    from collections import Counter
    layouts = [None] * len(expanded)
    content = []
    for i, (_, slide, items, _) in enumerate(expanded):
        if slide.get('type') in ('cover', 'final'):
            continue
        layouts[i] = natural_layout(slide, items)
        content.append(i)
    if len(content) < 3:
        return layouts
    counts = Counter(layouts[i] for i in content)
    dominant, dom_count = counts.most_common(1)[0]
    if dom_count / len(content) <= threshold:
        return layouts
    allowed = int(len(content) * threshold)
    alternates = [l for l in ('columns', 'timeline', 'editorial') if l != dominant]
    alt = 0
    for i in content:
        if counts[dominant] <= allowed:
            break
        if layouts[i] != dominant:
            continue
        if any(_units(item) > 72 for item in expanded[i][2]):
            continue  # 长要点页保持 editorial，避免挤坏
        layouts[i] = alternates[alt % len(alternates)]
        alt += 1
        counts[dominant] -= 1
    return layouts


def build_scenes(data, template='luxury', image_mode='none'):
    from style_layouts import design_scene, natural_layout, STYLES
    template = template if template in STYLES else 'luxury'
    expanded = []
    for source_index, slide in enumerate(data.get('slides', [])):
        content = slide.get('content', [])
        items = [str(item) for item in content] if isinstance(content, list) else [str(content)]
        # Split dense points into continuation pages rather than shrinking body copy.
        pieces = []
        for item in items:
            while _units(item) > 130:
                split = 1
                while split < len(item) and _units(item[:split]) < 125:
                    split += 1
                pieces.append(item[:split])
                item = item[split:]
            if item:
                pieces.append(item)
        per_page = 2 if any(_units(item) > 72 for item in pieces) else 4
        chunks = [pieces[i:i+per_page] for i in range(0, len(pieces), per_page)] or [[]]
        if slide.get('type') in ('cover', 'final'):
            chunks = [items]
        for part, chunk in enumerate(chunks):
            expanded.append((source_index, slide, chunk, part))

    layouts = _assign_layouts(expanded, natural_layout)
    scenes = []
    for index, ((source_index, slide, items, part), layout) in enumerate(zip(expanded, layouts)):
        draw_slide = slide if layout is None else dict(slide, layout=layout)
        scenes.append(design_scene(data, draw_slide, items, part, source_index, index, len(expanded), template, image_mode))
    return scenes


def render_scene(prs, scene, include_placeholder=True):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(scene['background'][1:])
    for el in scene['elements']:
        if not include_placeholder and el.get('role') == 'image-placeholder':
            continue
        if el['kind'] in ('rect', 'ellipse'):
            shape = slide.shapes.add_shape(9 if el['kind'] == 'ellipse' else 1, Inches(el['x']), Inches(el['y']), Inches(el['w']), Inches(el['h']))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(el['color'][1:])
            shape.line.fill.background()
        else:
            shape = slide.shapes.add_textbox(Inches(el['x']), Inches(el['y']), Inches(el['w']), Inches(el['h']))
            frame = shape.text_frame
            frame.word_wrap = False
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            for i, line in enumerate(el['text'].split('\n')):
                p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                p.text = line
                p.font.name = el.get('font', 'Microsoft YaHei')
                p.font.size = Pt(el['size'])
                p.font.bold = el['bold']
                p.font.color.rgb = RGBColor.from_string(el['color'][1:])
                p.line_spacing = 1.22
                p.space_after = Pt(0)
    return slide
