import base64
import io
import json
import os
import unittest
from unittest.mock import patch

from pptx import Presentation
from app import app, generate_default_content, outline_prompt
from generation_api import complete_slides
from generation_service import build_outline, validate_generation_options, validate_outline_request


class StreamTests(unittest.TestCase):
    def setUp(self):
        # 测试隔离：清除访问口令，避免 .env 中的 ACCESS_PASSWORD 让所有请求变 401。
        self._pw = patch.dict(os.environ, {'ACCESS_PASSWORD': ''})
        self._pw.start()
        self.addCleanup(self._pw.stop)
        self.client = app.test_client()
        self.outline = generate_default_content('流式测试', 4)

    def events(self, path, payload):
        response = self.client.post(path, json=payload)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        return [json.loads(line) for line in response.data.splitlines()]

    def test_prompt_preserves_constraints(self):
        brief = '管理者培训\n不要编造数据；只讨论医疗场景。'.ljust(200, '补')
        self.assertIn(brief, outline_prompt(brief, 8))
        events = self.events('/api/outline/stream', {'topic': brief, 'pages': 3, 'useAi': False})
        self.assertEqual(events[-1]['mode'], 'template')
        self.assertNotIn('\n', events[-1]['outline']['title'])
        self.assertEqual(len(events[-1]['outline']['slides']), 3)

    def test_ai_missing_key_does_not_silently_make_template(self):
        with patch('runtime_config.get_config', return_value={'DEEPSEEK_API_KEY': '', 'VISION_API_KEY': ''}):
            events = self.events('/api/outline/stream', {'topic': '主题', 'useAi': True})
        self.assertEqual(events[-1]['type'], 'error')
        self.assertIn('密钥', events[-1]['error'])

    def test_ai_chunked_outline_and_constraints(self):
        text = json.dumps(self.outline, ensure_ascii=False)
        class Upstream:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def raise_for_status(self): pass
            def iter_lines(self, **kwargs):
                for start in range(0, len(text), 7):
                    yield 'data: ' + json.dumps({'choices': [{'delta': {'content': text[start:start+7]}}]})
                yield 'data: ' + json.dumps({'choices': [{'delta': {}, 'finish_reason': 'stop'}]})
                yield 'data: [DONE]'
        with patch('runtime_config.get_config', return_value={'DEEPSEEK_API_KEY': 'test', 'VISION_API_KEY': ''}), patch('generation_service.requests.post', return_value=Upstream()) as post:
            response = self.client.post('/api/outline/stream', json={'topic': '只介绍医疗，不编造数据', 'useAi': True}, buffered=False)
            stream = iter(response.response)
            self.assertIn('stage', next(stream).decode())
            post.assert_not_called()  # Initial event reaches client before contacting upstream.
            events = [json.loads(line) for chunk in stream for line in chunk.splitlines()]
            self.assertEqual([e['slide'] for e in events if e['type']=='outline_slide'], self.outline['slides'])
            self.assertEqual(events[-1]['outline'], self.outline)
            self.assertIn('不编造数据', post.call_args.kwargs['json']['messages'][0]['content'])

    def test_partial_json_does_not_emit_unfinished_slide(self):
        text = '{"slides":[{"title":"含 } 和 \\\" 引号","type":"cover","content":"x"}, {"title":"未完'
        self.assertEqual(len(complete_slides(text)), 1)

    def test_ppt_stream_matches_export(self):
        response = self.client.post('/api/generate/stream', json={'outline': self.outline, 'imageMode':'none'}, buffered=False)
        iterator = iter(response.response)
        self.assertEqual(json.loads(next(iterator))['type'], 'stage')
        events = [json.loads(chunk) for chunk in iterator]
        scenes = [event['scene'] for event in events if event['type']=='slide']
        final = events[-1]
        self.assertEqual(final['type'], 'complete')
        ppt = Presentation(io.BytesIO(base64.b64decode(final['file'])))
        self.assertEqual(len(ppt.slides), len(scenes))
        for slide, scene in zip(ppt.slides, scenes):
            self.assertEqual([shape.text for shape in slide.shapes if shape.has_text_frame and shape.text],
                             [el['text'] for el in scene['elements'] if el['kind']=='text'])

    def test_search_failure_keeps_outline_and_never_calls_ai(self):
        with patch('pptx_export._generate_image_queries', return_value={}), patch('image_service._search_bing_images', return_value=[]), patch('image_service._ai_image') as ai:
            events = self.events('/api/generate/stream', {'outline':self.outline,'imageMode':'real'})
        self.assertEqual(events[-1]['code'], 'IMAGE_SEARCH_FAILED')
        self.assertEqual(events[-1]['outline'], self.outline)
        ai.assert_not_called()

    def test_bad_input_rejected_before_stream(self):
        for payload in ({'topic':' '}, {'topic':'x'*201}, {'topic':'x','pages':'bad'}, {'topic':'x','pages':31}):
            self.assertEqual(self.client.post('/api/outline/stream', json=payload).status_code, 400)
        for payload in ({'outline':{}}, {'outline':{'slides':[None]}}, {'outline':self.outline,'imageMode':'invalid'}):
            self.assertEqual(self.client.post('/api/generate/stream', json=payload).status_code, 400)

    def test_shared_request_validation_keeps_legacy_page_defaults(self):
        self.assertEqual(validate_outline_request({'topic':'主题'}, default_pages=5, clamp_pages=True)['pages'], 5)
        self.assertEqual(validate_outline_request({'topic':'主题','pages':99}, clamp_pages=True)['pages'], 30)
        with self.assertRaises(ValueError):
            validate_outline_request({'topic':'主题','pages':31})
        with self.assertRaises(ValueError):
            validate_outline_request({'topic':'x' * 201})

    def test_legacy_routes_share_outline_and_option_validation(self):
        self.assertEqual(self.client.post('/api/outline', json={'topic':'主题','pages':99}).status_code, 200)
        invalid_outline = {'title':'x','slides':[{'type':'unexpected','title':'x','content':''}]}
        self.assertFalse(self.client.post('/api/generate', json={'outline':invalid_outline}).json['success'])
        for path in ('/api/preview', '/api/generate'):
            response = self.client.post(path, json={'outline':self.outline,'template':'invalid'})
            self.assertEqual(response.status_code, 400)

    def test_build_outline_validates_generated_and_imported_data(self):
        import outline_service as backend
        self.assertEqual(build_outline(backend, '主题', 3)['slides'][0]['title'], '主题')
        with self.assertRaises(ValueError):
            build_outline(backend, '主题', 3, outline={'slides':[None]})
        self.assertEqual(validate_generation_options({'template':'modern','imageMode':'blank'})[:2],
                         ('modern', 'blank'))

    def test_ai_image_failure_is_explicit(self):
        with patch('pptx_export._generate_image_queries', return_value={}), patch('image_service._ai_image', return_value=None):
            events = self.events('/api/generate/stream', {'outline': self.outline,'imageMode':'ai'})
        self.assertEqual(events[-1]['type'], 'error')
        self.assertIn('AI 配图', events[-1]['error'])

    def test_output_limit_fails_before_oversized_event(self):
        oversized = iter([{'type':'stage','message':'准备'}, {'type':'complete','file':'x'*4_200_000}])
        # Use a generator so disconnect cleanup has close().
        def events():
            yield from oversized
        with patch('pptx_export.presentation_events', return_value=events()):
            result = self.events('/api/generate/stream', {'outline':self.outline})
        self.assertEqual(result[-1]['type'], 'error')
        self.assertIn('大小', result[-1]['error'])
        self.assertFalse(any(event['type']=='complete' for event in result))

    def test_truncated_ai_is_not_accepted(self):
        class Truncated:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def raise_for_status(self): pass
            def iter_lines(self, **kwargs):
                yield 'data: ' + json.dumps({'choices':[{'delta':{'content':'{"slides":['},'finish_reason':'length'}]})
        with patch('runtime_config.get_config', return_value={'DEEPSEEK_API_KEY': 'test', 'VISION_API_KEY': ''}), patch('generation_service.requests.post', return_value=Truncated()):
            result = self.events('/api/outline/stream', {'topic':'测试','useAi':True})
        self.assertEqual(result[-1]['type'], 'error')
        self.assertIn('未完成', result[-1]['error'])

    def test_budget_failure_releases_request_context(self):
        from generation_runtime import deadline
        with patch('pptx_export.check_budget', side_effect=ValueError('生成超时')):
            result = self.events('/api/generate/stream', {'outline':self.outline})
        self.assertEqual(result[-1]['type'], 'error')
        self.assertIsNone(deadline.get())


if __name__ == '__main__':
    unittest.main()
