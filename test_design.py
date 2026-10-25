import base64
import io
import os
import unittest
from unittest.mock import patch
from pptx import Presentation
from app import app, create_presentation, generate_default_content
from presentation_design import build_scenes


class DesignTests(unittest.TestCase):
    def setUp(self):
        # 测试隔离：清除访问口令，避免 .env 中的 ACCESS_PASSWORD 让所有请求变 401。
        self._pw = patch.dict(os.environ, {'ACCESS_PASSWORD': ''})
        self._pw.start()
        self.addCleanup(self._pw.stop)
        self.outline = {'title': '项目提案', 'slides': [
            {'type': 'cover', 'title': '项目提案', 'content': '设计与验证'},
            {'type': 'content', 'title': '发展历程', 'content': ['1956 起点', '2017 技术突破', '2026 实践应用']},
            {'type': 'content', 'title': '核心价值', 'content': ['提升效率', '可编辑输出', '统一设计']},
            {'type': 'final', 'title': '讨论与交流', 'content': 'Q&A'},
        ]}

    def test_export_matches_preview_and_is_editable(self):
        for theme in ('default', 'modern', 'creative', 'luxury'):
            scenes = build_scenes(self.outline, theme, 'none')
            _, raw = create_presentation(self.outline, theme, 'none')
            prs = Presentation(io.BytesIO(raw))
            self.assertEqual(len(prs.slides), len(scenes))
            for slide, scene in zip(prs.slides, scenes):
                actual = [s.text for s in slide.shapes if s.has_text_frame and s.text]
                expected = [el['text'] for el in scene['elements'] if el['kind'] == 'text']
                self.assertEqual(actual, expected)
                for shape in slide.shapes:
                    self.assertGreaterEqual(shape.left, 0)
                    self.assertGreaterEqual(shape.top, 0)
                    self.assertLessEqual(shape.left+shape.width, prs.slide_width+100)
                    self.assertLessEqual(shape.top+shape.height, prs.slide_height+100)

    def test_dense_content_continues_without_losing_text(self):
        points = ['长文本内容' * 60, *[f'论点 {i}' for i in range(10)]]
        scenes = build_scenes({'slides':[{'type':'content','title':'测试','content':points}]})
        self.assertGreater(len(scenes), 1)
        text = ''.join(el['text'].replace('\n','') for scene in scenes for el in scene['elements'] if el['kind']=='text')
        self.assertIn(points[0][:125], text)
        for point in points[1:]:
            self.assertIn(point, text)

    def test_layout_selection_and_single_page(self):
        scenes = build_scenes(self.outline)
        self.assertEqual(scenes[1]['layout'], 'timeline')
        self.assertEqual(scenes[2]['layout'], 'columns')
        self.assertEqual(len(generate_default_content('主题', 1)['slides']), 1)

    def test_image_does_not_remove_content(self):
        from PIL import Image
        buffer = io.BytesIO()
        Image.new('RGB', (400,300), 'green').save(buffer, format='PNG')
        with patch('pptx_export._generate_image_queries', return_value={}), patch('pptx_export._get_slide_image', return_value=buffer.getvalue()):
            _, raw = create_presentation(self.outline, 'luxury', 'real')
        slide = Presentation(io.BytesIO(raw)).slides[1]
        text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
        self.assertIn('1956', text)
        self.assertIn('2026', text)
        self.assertNotIn('IMAGE', text)
        self.assertTrue(any(s.shape_type == 13 for s in slide.shapes))

    def test_api_local_generation(self):
        client = app.test_client()
        result = client.post('/api/generate', json={'outline':self.outline,'template':'luxury','imageMode':'none','useAi':False}).get_json()
        self.assertTrue(result['success'])
        self.assertEqual(len(Presentation(io.BytesIO(base64.b64decode(result['file']))).slides), 4)
        preview = client.post('/api/preview', json={'outline':self.outline,'template':'luxury','imageMode':'none'})
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(len(preview.json['scenes']), 4)
        self.assertEqual(client.post('/api/preview', json={}).status_code, 400)

    def test_styles_differ_without_color(self):
        # A palette-only implementation must fail this regression check.
        for slide_index in (0, 1, 2):
            signatures = []
            for theme in ('default', 'modern', 'creative', 'luxury'):
                scene = build_scenes(self.outline, theme)[slide_index]
                signatures.append(tuple((el['kind'], el['x'], el['y'], el['w'], el['h'],
                                         el.get('size'), el.get('font'), el.get('bold'))
                                        for el in scene['elements']))
            self.assertEqual(len(set(signatures)), 4)

    def test_search_failure_requires_choice_and_never_calls_ai(self):
        import requests
        for candidates in ([], [{'url':'https://example.invalid/image', 'title':'test'}]):
            with patch('pptx_export._generate_image_queries', return_value={}), patch('image_service._search_bing_images', return_value=candidates), patch('image_service._download_image_bytes', return_value=None), patch('image_service._ai_image') as ai:
                response = app.test_client().post('/api/generate', json={'outline':self.outline, 'imageMode':'real'})
                self.assertEqual(response.status_code, 409)
                self.assertEqual(response.json['code'], 'IMAGE_SEARCH_FAILED')
                self.assertEqual(response.json['outline'], self.outline)
                self.assertEqual(response.json['sourceIndex'], 1)
                self.assertNotIn('file', response.json)
                ai.assert_not_called()
        with patch('pptx_export._generate_image_queries', return_value={}), patch('pptx_export._get_slide_image', side_effect=requests.Timeout()):
            response = app.test_client().post('/api/generate', json={'outline':self.outline, 'imageMode':'real'})
            self.assertEqual(response.json['code'], 'IMAGE_SEARCH_FAILED')

    def test_fallback_modes_do_not_search_again(self):
        from PIL import Image
        buffer = io.BytesIO()
        Image.new('RGB', (400, 300), 'green').save(buffer, format='PNG')
        for mode in ('blank', 'none', 'ai'):
            with patch('pptx_export._generate_image_queries', return_value={}), patch('image_service._search_bing_images') as search, patch('image_service._ai_image', return_value=buffer.getvalue()) as ai:
                response = app.test_client().post('/api/generate', json={'outline':self.outline, 'imageMode':mode})
                self.assertTrue(response.json['success'])
                search.assert_not_called()
                self.assertEqual(ai.call_count, 2 if mode=='ai' else 0)
                scenes = build_scenes(self.outline, image_mode=mode)
                self.assertEqual('image' in scenes[1], mode!='none')

    def test_all_layouts_preserve_content_and_bounds(self):
        for theme in ('default', 'modern', 'creative', 'luxury'):
            for layout in ('editorial', 'columns', 'timeline'):
                for mode in ('none', 'blank'):
                    for points in (['完整论点甲', '完整论点乙', '完整论点丙', '完整论点丁'], ['长文本内容'*26]*4):
                        scenes = build_scenes({'slides':[{'title':'标题', 'type':'content', 'content':points, 'layout':layout}]}, theme, mode)
                        text = ''.join(el['text'].replace('\n','') for scene in scenes for el in scene['elements'] if el['kind']=='text')
                        for point in points:
                            self.assertIn(point, text)
                        for scene in scenes:
                            for el in scene['elements']:
                                self.assertGreaterEqual(el['x'], 0)
                                self.assertGreaterEqual(el['y'], 0)
                                self.assertLessEqual(el['x']+el['w'], 13.333)
                                self.assertLessEqual(el['y']+el['h'], 7.5)
                                if el['kind']=='text':
                                    self.assertLessEqual(len(el['text'].split('\n'))*el['size']*1.22, el['h']*72+.01)


if __name__ == '__main__':
    unittest.main()
