"""Render shared scenes and emit in-memory PPTX generation events."""
import base64
import contextvars
import io
import logging
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

import requests
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from generation_runtime import check_budget
from image_service import _generate_image_queries, _get_slide_image
from presentation_design import build_scenes, render_scene

logger = logging.getLogger(__name__)


class ImageSearchFailed(Exception):
    def __init__(self, source_index):
        super().__init__('联网搜索未能取得可用图片，请选择后续处理方式。')
        self.source_index = source_index


def _compact_image(image_bytes):
    from PIL import Image, ImageOps
    with Image.open(io.BytesIO(image_bytes)) as picture:
        picture = ImageOps.exif_transpose(picture).convert('RGB')
        picture.thumbnail((1280, 960))
        compact = io.BytesIO()
        picture.save(compact, format='JPEG', quality=82, optimize=True)
        return compact.getvalue()


def _prepare_image(data, topic, image_mode, idx, query, use_vision):
    """取一页配图并压缩；异常按生成语义抛出（供并发预取调用）。"""
    try:
        image = _get_slide_image(data['slides'][idx], topic, image_mode, query=query, use_vision=use_vision)
    except requests.RequestException as error:
        if image_mode == 'real':
            raise ImageSearchFailed(idx) from error
        raise
    if image_mode == 'real' and not image:
        raise ImageSearchFailed(idx)
    if image_mode == 'ai' and not image:
        raise ValueError('AI 配图未成功，请重试或返回设置选择留空 / 纯排版。')
    return _compact_image(image)


def presentation_events(data, template='default', image_mode='blank', use_vision=False):
    """创建PPT文件"""
    prs = Presentation()

    # 设置幻灯片尺寸为16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    topic = data.get('title', '')
    yield {'type': 'stage', 'message': '正在准备排版与配图'}
    queries = _generate_image_queries(data) if image_mode in ('ai', 'real') else {}
    images = {}
    scenes = build_scenes(data, template, image_mode)

    # 并发预取所有需要的配图（下载是 IO 密集，串行是主要瓶颈）
    if image_mode in ('ai', 'real'):
        needed = list(dict.fromkeys(scene['sourceIndex'] for scene in scenes if 'image' in scene))
        if needed:
            yield {'type': 'stage', 'message': f'正在并发下载配图（{len(needed)} 页）'}
            ctx = contextvars.copy_context()
            with ThreadPoolExecutor(max_workers=min(8, len(needed))) as pool:
                futures = {idx: pool.submit(ctx.copy().run, _prepare_image, data, topic, image_mode, idx,
                                            queries.get(idx), use_vision) for idx in needed}
                for idx, fut in futures.items():
                    images[idx] = fut.result()

    for number, scene in enumerate(scenes):
        check_budget()
        yield {'type': 'stage', 'message': f'正在制作第 {number + 1} / {len(scenes)} 页'}
        image = images.get(scene['sourceIndex']) if 'image' in scene else None
        slide = render_scene(prs, scene, include_placeholder=image is None)
        if image:
            _add_image(slide, image, left=scene['image']['x'], top=scene['image']['y'],
                       width=scene['image']['w'], height=scene['image']['h'])
        preview_scene = dict(scene)
        if image:
            # 浏览器使用和 PPTX 相同的图片字节，避免仅显示图片占位框。
            from PIL import Image
            with Image.open(io.BytesIO(image)) as picture:
                preview_scene['imageData'] = 'data:' + Image.MIME.get(picture.format, 'image/png') + ';base64,' + base64.b64encode(image).decode('ascii')
        yield {'type': 'slide', 'index': number, 'total': len(scenes), 'scene': preview_scene}

    # 生成文件名（在内存中生成，适配 Vercel 只读文件系统，无需落盘）
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"PPT_{timestamp}_{uuid.uuid4().hex[:8]}.pptx"
    buf = io.BytesIO()
    prs.save(buf)
    check_budget()
    if buf.tell() > 3 * 1024 * 1024:
        raise ValueError('PPTX 文件较大，请减少配图或拆分文稿后重试。')
    yield {'type': 'complete', 'success': True, 'filename': filename,
           'file': base64.b64encode(buf.getvalue()).decode('ascii'), 'outline': data}


def create_presentation(data, template='default', image_mode='blank', use_vision=False):
    for event in presentation_events(data, template, image_mode, use_vision):
        if event['type'] == 'complete':
            return event['filename'], base64.b64decode(event['file'])


def _add_image(slide, image_bytes, left, top, width, height):
    """等比插入图片并居中，失败时退回占位框"""
    try:
        pic = slide.shapes.add_picture(io.BytesIO(image_bytes), Inches(left), Inches(top), width=Inches(width))
        if pic.height.inches > height:
            scale = height / pic.height.inches
            pic.width = Inches(width * scale)
            pic.height = Inches(height)
        pic.left = Inches(left + (width - pic.width.inches) / 2)
        pic.top = Inches(top + (height - pic.height.inches) / 2)
    except Exception as e:
        logger.warning("插入图片失败: %s", e)
        _add_image_placeholder(slide, left, top, width, height)


def _add_image_placeholder(slide, left, top, width, height):
    """留空占位框，方便用户后续自定义"""
    rect = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    rect.fill.background()
    rect.line.color.rgb = RGBColor(200, 200, 200)
    rect.line.width = Pt(1)
    tf = rect.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = '此处插入图片'
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(160, 160, 160)
